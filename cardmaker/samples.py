"""生成示例数据（示例 xlsx + 示例配图），供界面「生成示例数据」按钮和命令行使用。"""
from __future__ import annotations

import os
from typing import List, Tuple

from PIL import Image, ImageDraw

DATA = [
    {
        "标题": "炎龙战士", "副标题": "FIRE DRAGON WARRIOR",
        "描述": "由火山深处苏醒的古老守护者，双翼燃着永不熄灭的烈焰。\n攻击时附带灼烧效果，持续 3 回合。",
        "页脚": "NO.001 / 稀有度 ★★★★★", "图片": "dragon.png",
        "标题.color": "#B22222", "副标题.color": "#E07A2F", "描述.size": "26",
    },
    {
        "标题": "深海使者", "副标题": "DEEP SEA ENVOY",
        "描述": "来自幽暗海沟的沉默使者，能操控洋流改变战场地形。\n技能：潮汐之力，每回合恢复 2 点能量。",
        "页脚": "NO.002 / 稀有度 ★★★★", "图片": "ocean.png",
        "标题.color": "#1B5E8C", "副标题.color": "#2E86AB", "描述.size": "26",
    },
    {
        "标题": "林间游侠", "副标题": "FOREST RANGER",
        "描述": "在千年古林中长大的猎手，箭无虚发。\n技能：疾风步，行动后额外获得一次移动机会。",
        "页脚": "NO.003 / 稀有度 ★★★★", "图片": "forest.png",
        "标题.color": "#2E7D32", "副标题.color": "#4CAF50", "描述.size": "26",
    },
    {
        "标题": "星穹法师", "副标题": "STARFORGE MAGE",
        "描述": "观星者一族的最后传人，能以星辰为引编织法术。\n技能：陨星召唤，对全体敌人造成大量伤害。",
        "页脚": "NO.004 / 稀有度 ★★★★★★", "图片": "star.png",
        "标题.color": "#4A3C8C", "副标题.color": "#7E57C2", "描述.size": "26",
    },
    {
        "标题": "沙漠之影", "副标题": "DESERT SHADE",
        "描述": "隐匿于黄沙之间的刺客，无人见过其真容。\n技能：幻影突袭，攻击后立即可再次行动。",
        "页脚": "NO.005 / 稀有度 ★★★★★", "图片": "desert.png",
        "标题.color": "#C07A1B", "副标题.color": "#D9A441", "描述.size": "24",
    },
    {
        "标题": "极地守卫", "副标题": "POLAR GUARDIAN",
        "描述": "霜之堡垒的永恒哨兵，躯体由万年寒冰凝成。\n被动：寒霜护甲，受到物理伤害降低 40%。",
        "页脚": "NO.006 / 稀有度 ★★★★", "图片": "polar.png",
        "标题.color": "#1E6B7A", "副标题.color": "#4FB3C4", "描述.size": "26",
    },
]

_COLORS = [
    ((216, 78, 62), (255, 205, 150)),
    ((30, 96, 140), (150, 220, 245)),
    ((46, 125, 50), (186, 232, 160)),
    ((74, 60, 140), (208, 190, 250)),
    ((192, 122, 27), (250, 224, 165)),
    ((30, 107, 122), (176, 232, 240)),
]
_NAMES = ["dragon", "ocean", "forest", "star", "desert", "polar"]


def make_sample_image(path: str, c1, c2, size: Tuple[int, int] = (640, 640)) -> None:
    img = Image.new("RGB", size, c1)
    d = ImageDraw.Draw(img)
    for y in range(size[1]):
        t = y / max(1, size[1] - 1)
        col = tuple(int(c1[i] + (c2[i] - c1[i]) * t) for i in range(3))
        d.line([(0, y), (size[0], y)], fill=col)
    for i in range(6):
        r = 40 + i * 46
        d.ellipse([size[0] // 2 - r, size[1] // 2 - r,
                   size[0] // 2 + r, size[1] // 2 + r],
                  outline=(255, 255, 255), width=3)
    d.ellipse([size[0] // 2 - 70, size[1] // 2 - 70,
               size[0] // 2 + 70, size[1] // 2 + 70], fill=(255, 255, 255))
    img.save(path)


def generate(base_dir: str) -> Tuple[str, str]:
    """在 base_dir 下生成 samples/示例卡片.xlsx 与 samples/images/*.png。"""
    sample_dir = os.path.join(base_dir, "samples")
    img_dir = os.path.join(sample_dir, "images")
    os.makedirs(img_dir, exist_ok=True)
    for i, n in enumerate(_NAMES):
        make_sample_image(os.path.join(img_dir, f"{n}.png"), _COLORS[i][0], _COLORS[i][1])

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "卡片数据"
    headers: List[str] = list(DATA[0].keys())
    ws.append(headers)
    for row in DATA:
        vals = []
        for h in headers:
            v = row.get(h, "")
            if h == "图片":
                v = os.path.join("samples", "images", str(v))
            vals.append(v)
        ws.append(vals)

    head_font = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4472C4")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, h in enumerate(headers, start=1):
        width = 52 if h == "描述" else max(13, min(40, len(h) * 2 + 10))
        col_letter = chr(64 + i) if i <= 26 else "A"
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    out = os.path.join(sample_dir, "示例卡片.xlsx")
    wb.save(out)
    return out, img_dir
